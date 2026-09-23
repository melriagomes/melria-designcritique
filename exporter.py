"""Exports a finished critique (Markdown reply + annotated image) as a Word
document or a PDF.

The critique Markdown uses a small fixed subset — `##`/`###` headings, `-`
and `1.` lists, blank-line paragraphs, and `**bold**` — the same subset
index.html's renderMarkdown() displays. `parse_markdown` turns it into
blocks once; `to_docx` and `to_pdf` render those blocks.

The PDF is printed by headless Chromium (already used for webpage
screenshots) from HTML built here with every piece of text escaped, with
JavaScript disabled and all network requests blocked, so nothing in a
critique can make the server fetch anything.
"""
import base64
import binascii
import html
import io
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from playwright.sync_api import sync_playwright

TITLE = "Design Critique"
DATA_URL_RE = re.compile(r"^data:image/(png|jpeg);base64,([A-Za-z0-9+/=]+)$")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
BULLET_RE = re.compile(r"^[-*]\s+")
NUMBERED_RE = re.compile(r"^\d+\.\s+")
ACCENT = RGBColor(0x6F, 0x00, 0xFF)


def decode_image_data_url(data_url):
    """Return (image_bytes, mime) for a PNG/JPEG data URL, or (None, None)
    when it's missing or isn't a real image."""
    match = DATA_URL_RE.match((data_url or "").strip())
    if not match:
        return None, None
    try:
        data = base64.b64decode(match.group(2), validate=True)
        with Image.open(io.BytesIO(data)) as img:
            img.verify()
    except (binascii.Error, ValueError, OSError):
        return None, None
    return data, f"image/{match.group(1)}"


def parse_markdown(md):
    """Parse critique Markdown into blocks: ("h2"|"h3"|"p", text) or
    ("ul"|"ol", [item_text, ...])."""
    blocks = []
    paragraph = []
    current_list = None

    def flush_paragraph():
        if paragraph:
            blocks.append(("p", " ".join(paragraph)))
            paragraph.clear()

    def close_list():
        nonlocal current_list
        current_list = None

    for raw in (md or "").split("\n"):
        line = raw.strip()
        if not line:
            flush_paragraph()
            close_list()
        elif line.startswith("### "):
            flush_paragraph()
            close_list()
            blocks.append(("h3", line[4:]))
        elif line.startswith("## "):
            flush_paragraph()
            close_list()
            blocks.append(("h2", line[3:]))
        elif BULLET_RE.match(line) or NUMBERED_RE.match(line):
            flush_paragraph()
            kind = "ul" if BULLET_RE.match(line) else "ol"
            text = (BULLET_RE if kind == "ul" else NUMBERED_RE).sub("", line)
            if current_list is None or current_list[0] != kind:
                current_list = (kind, [])
                blocks.append(current_list)
            current_list[1].append(text)
        else:
            close_list()
            paragraph.append(line)
    flush_paragraph()
    return blocks


def inline_runs(text):
    """Split text into [(segment, is_bold), ...] on **bold** markers."""
    runs = []
    pos = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > pos:
            runs.append((text[pos:match.start()], False))
        runs.append((match.group(1), True))
        pos = match.end()
    if pos < len(text):
        runs.append((text[pos:], False))
    return runs


def to_docx(md, image_bytes=None):
    """Render the critique as .docx bytes."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(TITLE, level=0)

    if image_bytes:
        doc.add_picture(io.BytesIO(image_bytes), width=Inches(6.0))

    def add_runs(paragraph, text):
        for segment, bold in inline_runs(text):
            run = paragraph.add_run(segment)
            if bold:
                run.bold = True
                run.font.color.rgb = ACCENT

    for kind, content in parse_markdown(md):
        if kind == "h2":
            doc.add_heading(content.replace("**", ""), level=1)
        elif kind == "h3":
            doc.add_heading(content.replace("**", ""), level=2)
        elif kind in ("ul", "ol"):
            style = "List Bullet" if kind == "ul" else "List Number"
            for item in content:
                add_runs(doc.add_paragraph(style=style), item)
        else:
            add_runs(doc.add_paragraph(), content)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _inline_html(text):
    return "".join(
        f"<strong>{html.escape(seg)}</strong>" if bold else html.escape(seg)
        for seg, bold in inline_runs(text)
    )


def to_html(md, image_bytes=None, mime=None):
    """Render the critique as a standalone, print-styled HTML document."""
    parts = [f"<h1>{html.escape(TITLE)}</h1>"]
    if image_bytes:
        b64 = base64.b64encode(image_bytes).decode("ascii")
        parts.append(f'<img src="data:{mime};base64,{b64}" alt="Annotated design with numbered callouts">')
    for kind, content in parse_markdown(md):
        if kind in ("h2", "h3"):
            parts.append(f"<{kind}>{_inline_html(content)}</{kind}>")
        elif kind in ("ul", "ol"):
            items = "".join(f"<li>{_inline_html(item)}</li>" for item in content)
            parts.append(f"<{kind}>{items}</{kind}>")
        else:
            parts.append(f"<p>{_inline_html(content)}</p>")

    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="UTF-8"><title>{html.escape(TITLE)}</title>
<style>
  body {{ font-family: -apple-system, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
         color: #3B0270; font-size: 11pt; line-height: 1.55; margin: 0; }}
  h1 {{ font-size: 20pt; margin: 0 0 12pt; }}
  h2 {{ font-size: 14pt; margin: 18pt 0 6pt; }}
  h3 {{ font-size: 12pt; margin: 14pt 0 4pt; color: #6F00FF; }}
  h2, h3 {{ break-after: avoid; }}
  p {{ margin: 0 0 6pt; }}
  ul, ol {{ margin: 0 0 6pt; padding-left: 18pt; }}
  li {{ margin-bottom: 3pt; break-inside: avoid; }}
  strong {{ color: #6F00FF; }}
  img {{ display: block; box-sizing: border-box; max-width: 100%; border: 1px solid #F1D6F8; border-radius: 6pt; margin: 0 0 12pt; }}
</style></head><body>{"".join(parts)}</body></html>"""


def to_pdf(md, image_bytes=None, mime=None):
    """Render the critique as PDF bytes via headless Chromium."""
    document = to_html(md, image_bytes, mime)
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        try:
            context = browser.new_context(java_script_enabled=False)
            page = context.new_page()
            # Everything the document needs is inline; refuse any fetch.
            page.route("**/*", lambda route: route.abort())
            page.set_content(document, wait_until="load")
            return page.pdf(
                format="A4",
                print_background=True,
                margin={"top": "18mm", "bottom": "18mm", "left": "16mm", "right": "16mm"},
            )
        finally:
            browser.close()
